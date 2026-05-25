#!/usr/bin/env python3
"""测试 section break 保护和深度格式拷贝"""
import sys
import os
import tempfile
import shutil

# 添加 scripts 目录到 path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'scripts'))

import docx
from docx.oxml.ns import qn
from docx.shared import Pt, Emu


def test_remove_empty_paragraphs_preserves_sectPr():
    """测试 remove_empty_paragraphs 不会删除包含 pPr/sectPr 的空段落"""
    from fix_docx_template import remove_empty_paragraphs

    doc = docx.Document()

    # 添加一个普通段落
    doc.add_paragraph('Hello')

    # 添加一个空段落，但 pPr 内含 sectPr（模拟分节符）
    p_empty_with_sect = doc.add_paragraph('')
    pPr = p_empty_with_sect._element.get_or_add_pPr()
    sectPr = docx.oxml.OxmlElement('w:sectPr')
    sectPr.set(qn('w:pgSz'), '')  # minimal sectPr
    pPr.append(sectPr)

    # 添加另一个普通段落
    doc.add_paragraph('World')

    # 添加一个纯空段落（应被删除）
    doc.add_paragraph('')

    # 记录删除前的段落数
    before_count = len(doc.paragraphs)

    # 执行
    remove_empty_paragraphs(doc)

    after_count = len(doc.paragraphs)

    # 纯空段落应被删除（1个），含 sectPr 的应保留
    assert after_count == before_count - 1, f"Expected {before_count - 1} paragraphs, got {after_count}"

    # 确认含 sectPr 的段落仍然存在
    remaining_texts = [p.text for p in doc.paragraphs]
    assert '' in remaining_texts, "Empty paragraph with sectPr should be preserved"

    print("✓ test_remove_empty_paragraphs_preserves_sectPr PASSED")


def test_remove_empty_paragraphs_preserves_images():
    """测试 remove_empty_paragraphs 不会删除包含图片的空段落"""
    from fix_docx_template import remove_empty_paragraphs

    doc = docx.Document()
    doc.add_paragraph('Before')

    # 添加一个空段落但包含 drawing 元素
    p_img = doc.add_paragraph('')
    run = p_img.add_run()
    drawing = docx.oxml.OxmlElement('w:drawing')
    run._element.append(drawing)

    doc.add_paragraph('After')
    doc.add_paragraph('')  # 纯空段落

    before = len(doc.paragraphs)
    remove_empty_paragraphs(doc)
    after = len(doc.paragraphs)

    assert after == before - 1, f"Expected {before - 1}, got {after}"
    print("✓ test_remove_empty_paragraphs_preserves_images PASSED")


def test_copy_format_deep_basic():
    """测试深度格式拷贝的基本功能"""
    # 创建临时目录
    tmpdir = tempfile.mkdtemp()
    try:
        # 创建模板文档
        tmpl_path = os.path.join(tmpdir, 'template.docx')
        tmpl = docx.Document()

        # 段落1: 宋体 24pt 加粗 居中
        p1 = tmpl.add_paragraph('Title')
        p1.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        run1 = p1.runs[0]
        run1.font.name = 'SimSun'
        run1.font.size = Pt(24)
        run1.font.bold = True

        # 段落2: 楷体 14pt
        p2 = tmpl.add_paragraph('Body text')
        run2 = p2.runs[0]
        run2.font.name = 'KaiTi'
        run2.font.size = Pt(14)

        tmpl.save(tmpl_path)

        # 创建目标文档（相同结构，不同格式）
        tgt_path = os.path.join(tmpdir, 'target.docx')
        tgt = docx.Document()

        p1t = tgt.add_paragraph('Title')
        run1t = p1t.runs[0]
        run1t.font.name = 'Arial'
        run1t.font.size = Pt(10)

        p2t = tgt.add_paragraph('Body text')
        run2t = p2t.runs[0]
        run2t.font.name = 'Arial'
        run2t.font.size = Pt(10)

        tgt.save(tgt_path)

        # 执行深度拷贝
        sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'scripts'))
        from copy_format_deep import copy_format
        stats = copy_format(tmpl_path, tgt_path)

        # 验证
        assert stats['paragraphs'] == 2
        assert stats['tables'] == 0

        result = docx.Document(tgt_path)

        # 验证段落1格式
        p1r = result.paragraphs[0]
        assert p1r.alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        assert p1r.runs[0].font.size == Pt(24)
        assert p1r.runs[0].font.bold == True

        # 验证段落2格式
        p2r = result.paragraphs[1]
        assert p2r.runs[0].font.size == Pt(14)

        # 验证文字内容未被修改
        assert p1r.text == 'Title'
        assert p2r.text == 'Body text'

        print("✓ test_copy_format_deep_basic PASSED")

    finally:
        shutil.rmtree(tmpdir)


def test_copy_format_deep_preserves_content():
    """测试深度拷贝不修改文字内容"""
    tmpdir = tempfile.mkdtemp()
    try:
        tmpl_path = os.path.join(tmpdir, 'template.docx')
        tmpl = docx.Document()
        tmpl.add_paragraph('Template Title')
        tmpl.add_paragraph('Template Body')
        tmpl.save(tmpl_path)

        tgt_path = os.path.join(tmpdir, 'target.docx')
        tgt = docx.Document()
        tgt.add_paragraph('Target Title')  # 不同的文字
        tgt.add_paragraph('Target Body')
        tgt.save(tgt_path)

        from copy_format_deep import copy_format
        copy_format(tmpl_path, tgt_path)

        result = docx.Document(tgt_path)
        # 文字内容应保持不变
        assert result.paragraphs[0].text == 'Target Title'
        assert result.paragraphs[1].text == 'Target Body'

        print("✓ test_copy_format_deep_preserves_content PASSED")

    finally:
        shutil.rmtree(tmpdir)


def test_classify_cover_title_not_requiring_first_para():
    """测试封面标题匹配不要求 i==0"""
    from fix_docx_template import classify_and_format

    doc = docx.Document()

    # 模拟 "附件6" 开头的文档：前面有非《》的段落
    doc.add_paragraph('附件6')
    doc.add_paragraph('河南林业职业学院')
    doc.add_paragraph('《测试课程》')
    doc.add_paragraph('课程标准')

    classify_and_format(doc)

    # 验证 "《测试课程》" 段落（i=2）被应用了封面标题格式
    # 封面标题格式设置 first_line_indent = None
    cover_para = doc.paragraphs[2]
    assert cover_para.paragraph_format.first_line_indent is None, \
        f"Cover title should have no indent, got {cover_para.paragraph_format.first_line_indent}"

    print("✓ test_classify_cover_title_not_requiring_first_para PASSED")


if __name__ == '__main__':
    test_remove_empty_paragraphs_preserves_sectPr()
    test_remove_empty_paragraphs_preserves_images()
    test_copy_format_deep_basic()
    test_copy_format_deep_preserves_content()
    test_classify_cover_title_not_requiring_first_para()
    print()
    print("=" * 50)
    print("All tests passed!")
