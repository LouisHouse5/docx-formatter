#!/usr/bin/env python3
"""
深度格式拷贝：按段落索引一对一从模板复制格式到目标文档。
适用于两份文档结构一致的格式对齐场景。

与 fix_docx_template.py 的区别：
- fix_docx_template: 基于内容规则匹配（classify_and_format），适用于结构不同的文档
- copy_format_deep:  按索引直接拷贝格式，适用于结构一致的文档，零规则偏差

使用方法:
  python3 copy_format_deep.py <模板.docx> <目标.docx>
"""
import sys
import copy
import docx
from docx.oxml.ns import qn


def _deep_sync_pPr(src_para, dst_para):
    """深度同步段落格式：直接替换 pPr 元素"""
    src_pPr = src_para._element.find(qn('w:pPr'))
    dst_pPr = dst_para._element.find(qn('w:pPr'))

    if src_pPr is not None:
        new_pPr = copy.deepcopy(src_pPr)
        if dst_pPr is not None:
            dst_para._element.replace(dst_pPr, new_pPr)
        else:
            dst_para._element.insert(0, new_pPr)
    elif dst_pPr is not None:
        dst_para._element.remove(dst_pPr)


def _deep_sync_rPr(src_para, dst_para):
    """深度同步 run 格式：逐 run 替换 rPr 元素"""
    src_runs = src_para.runs
    dst_runs = dst_para.runs

    for sr, dr in zip(src_runs, dst_runs):
        src_rPr = sr._element.find(qn('w:rPr'))
        if src_rPr is not None:
            new_rPr = copy.deepcopy(src_rPr)
            old_rPr = dr._element.find(qn('w:rPr'))
            if old_rPr is not None:
                dr._element.replace(old_rPr, new_rPr)
            else:
                dr._element.insert(0, new_rPr)


def _copy_cell_format(src_cell, dst_cell):
    """复制单元格内段落格式"""
    for sp, dp in zip(src_cell.paragraphs, dst_cell.paragraphs):
        _deep_sync_pPr(sp, dp)
        _deep_sync_rPr(sp, dp)


def _copy_table_format(src_table, dst_table):
    """复制表格格式（tblPr、tblGrid、行属性、单元格属性+内容格式）"""
    src_tbl = src_table._tbl
    dst_tbl = dst_table._tbl

    # 1. 表格属性 (tblPr)
    src_tblPr = src_tbl.find(qn('w:tblPr'))
    if src_tblPr is not None:
        new_tblPr = copy.deepcopy(src_tblPr)
        old_tblPr = dst_tbl.find(qn('w:tblPr'))
        if old_tblPr is not None:
            dst_tbl.replace(old_tblPr, new_tblPr)
        else:
            dst_tbl.insert(0, new_tblPr)

    # 2. 表格网格 (tblGrid)
    src_grid = src_tbl.find(qn('w:tblGrid'))
    if src_grid is not None:
        new_grid = copy.deepcopy(src_grid)
        old_grid = dst_tbl.find(qn('w:tblGrid'))
        if old_grid is not None:
            dst_tbl.replace(old_grid, new_grid)
        else:
            tblPr = dst_tbl.find(qn('w:tblPr'))
            if tblPr is not None:
                tblPr.addnext(new_grid)

    # 3. 行属性 + 单元格格式
    for sr, dr in zip(src_table.rows, dst_table.rows):
        # 行属性 (trPr)
        src_trPr = sr._tr.find(qn('w:trPr'))
        if src_trPr is not None:
            new_trPr = copy.deepcopy(src_trPr)
            old_trPr = dr._tr.find(qn('w:trPr'))
            if old_trPr is not None:
                dr._tr.replace(old_trPr, new_trPr)
            else:
                dr._tr.insert(0, new_trPr)

        # 单元格属性 (tcPr) + 内容格式
        for sc, dc in zip(sr.cells, dr.cells):
            src_tcPr = sc._tc.find(qn('w:tcPr'))
            if src_tcPr is not None:
                new_tcPr = copy.deepcopy(src_tcPr)
                old_tcPr = dc._tc.find(qn('w:tcPr'))
                if old_tcPr is not None:
                    dc._tc.replace(old_tcPr, new_tcPr)
                else:
                    dc._tc.insert(0, new_tcPr)

            _copy_cell_format(sc, dc)


def _copy_section_settings(tmpl_doc, tgt_doc):
    """同步所有 section 的页面设置"""
    tmpl_sections = tmpl_doc.sections
    tgt_sections = tgt_doc.sections
    count = min(len(tmpl_sections), len(tgt_sections))

    for i in range(count):
        ts = tmpl_sections[i]
        s = tgt_sections[i]
        s.page_width = ts.page_width
        s.page_height = ts.page_height
        s.orientation = ts.orientation
        s.top_margin = ts.top_margin
        s.bottom_margin = ts.bottom_margin
        s.left_margin = ts.left_margin
        s.right_margin = ts.right_margin
        s.header_distance = ts.header_distance
        s.footer_distance = ts.footer_distance
        s.different_first_page_header_footer = ts.different_first_page_header_footer

    # 目标 section 更多时，用最后一个模板 section 的设置
    if len(tgt_sections) > count:
        last_tmpl = tmpl_sections[-1]
        for i in range(count, len(tgt_sections)):
            s = tgt_sections[i]
            s.page_width = last_tmpl.page_width
            s.page_height = last_tmpl.page_height
            s.orientation = last_tmpl.orientation
            s.top_margin = last_tmpl.top_margin
            s.bottom_margin = last_tmpl.bottom_margin
            s.left_margin = last_tmpl.left_margin
            s.right_margin = last_tmpl.right_margin
            s.header_distance = last_tmpl.header_distance
            s.footer_distance = last_tmpl.footer_distance
            s.different_first_page_header_footer = last_tmpl.different_first_page_header_footer

    print(f"  已同步 {len(tgt_sections)} 个 section 的页面设置")
    return count


def copy_format(template_path, target_path):
    """
    主函数：从模板深度复制格式到目标文档。

    返回 dict 包含同步统计信息。
    """
    tmpl_doc = docx.Document(template_path)
    tgt_doc = docx.Document(target_path)

    stats = {}

    # 1. 逐段落同步格式
    tmpl_paras = tmpl_doc.paragraphs
    tgt_paras = tgt_doc.paragraphs
    para_count = min(len(tmpl_paras), len(tgt_paras))

    for i in range(para_count):
        _deep_sync_pPr(tmpl_paras[i], tgt_paras[i])
        _deep_sync_rPr(tmpl_paras[i], tgt_paras[i])

    stats['paragraphs'] = para_count
    stats['paragraphs_extra'] = max(0, len(tgt_paras) - len(tmpl_paras))
    print(f"  已同步 {para_count} 个段落格式"
          + (f" (目标多出 {stats['paragraphs_extra']} 个未处理)" if stats['paragraphs_extra'] else ""))

    # 2. 逐表格同步格式
    tmpl_tables = tmpl_doc.tables
    tgt_tables = tgt_doc.tables
    table_count = min(len(tmpl_tables), len(tgt_tables))

    for i in range(table_count):
        _copy_table_format(tmpl_tables[i], tgt_tables[i])

    stats['tables'] = table_count
    print(f"  已同步 {table_count} 个表格格式")

    # 3. 同步 section 设置
    stats['sections'] = _copy_section_settings(tmpl_doc, tgt_doc)

    # 4. 保存
    tgt_doc.save(target_path)
    stats['saved'] = True

    return stats


def main():
    if len(sys.argv) < 3:
        print("用法: python3 copy_format_deep.py <模板.docx> <目标.docx>")
        print()
        print("深度格式拷贝：按段落索引一对一从模板复制格式到目标文档。")
        print("适用于两份文档结构一致的格式对齐场景。")
        print("只复制格式，不修改文字内容。")
        sys.exit(1)

    template_path = sys.argv[1]
    target_path = sys.argv[2]

    print(f"模板: {template_path}")
    print(f"目标: {target_path}")
    print("-" * 60)

    stats = copy_format(template_path, target_path)

    print("-" * 60)
    print(f"✓ 深度格式同步完成！已保存: {target_path}")
    print(f"  段落: {stats['paragraphs']}, 表格: {stats['tables']}, Section: {stats['sections']}")


if __name__ == '__main__':
    main()
