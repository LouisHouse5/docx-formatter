#!/usr/bin/env python3
"""
修复目录页码：将 ** 占位符替换为实际页码。
通过文档结构分析估算每个标题所在页码。
"""
import os
import math
import sys
from docx import Document
from docx.shared import Pt, Emu, Length
from docx.enum.text import WD_LINE_SPACING
from utils import has_toc_field


BASE_DIR = '/Users/chutianshu/Documents/school-work/办公文件/26 级移动互联课程标准'
TOC_RANGE = range(20, 39)

NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _line_height_pt(font_pt, pf):
    """计算段落行高（pt），正确处理 EMU vs 倍数"""
    ls = pf.line_spacing
    rule = pf.line_spacing_rule

    # 有明确的 rule 时
    if rule is not None:
        if rule == WD_LINE_SPACING.SINGLE:
            return font_pt * 1.0
        elif rule == WD_LINE_SPACING.ONE_POINT_FIVE:
            return font_pt * 1.5
        elif rule == WD_LINE_SPACING.DOUBLE:
            return font_pt * 2.0
        elif rule == WD_LINE_SPACING.MULTIPLE:
            # ls 是倍数 (float)
            return font_pt * ls
        elif rule in (WD_LINE_SPACING.EXACTLY, WD_LINE_SPACING.AT_LEAST):
            # ls 是 Length/EMU
            return ls / 12700

    # rule 为 None，看 ls 的类型判断
    if ls is None:
        # 继承样式，默认 1.5 倍
        return font_pt * 1.5

    if isinstance(ls, float):
        # float → 倍数
        return font_pt * ls
    elif isinstance(ls, (int, Length)):
        # int/Length → EMU
        return ls / 12700

    return font_pt * 1.5


def _get_font_pt(para, default=Pt(12)):
    """获取段落字体大小（pt）"""
    for r in para.runs:
        if r.font.size:
            return r.font.size / 12700
    return default / 12700


def _para_height_pt(para, usable_w_pt, default_font_pt=12):
    """估算段落高度（pt）"""
    pf = para.paragraph_format
    font_pt = _get_font_pt(para, Pt(default_font_pt))
    line_h = _line_height_pt(font_pt, pf)

    sp_before = (pf.space_before or 0) / 12700
    sp_after = (pf.space_after or 0) / 12700

    text = ''.join([r.text or '' for r in para.runs])

    if not text.strip():
        return line_h + sp_before + sp_after

    # 估算行数：中文1字符=1字号宽，半角=0.5字号宽
    char_count = sum(1 if ord(ch) > 127 else 0.5 for ch in text)
    chars_per_line = max(1, usable_w_pt / font_pt)
    lines = max(1, math.ceil(char_count / chars_per_line))

    return lines * line_h + sp_before + sp_after


def _table_height_pt(table):
    """估算表格高度（pt）"""
    # 每行高度约 25-35pt，取决于内容
    rows = len(table.rows)
    if rows == 0:
        return 0
    # 检查是否有显式行高
    total = 0
    for row in table.rows:
        tr = row._tr
        trPr = tr.find(f'{NS}trPr')
        if trPr is not None:
            trHeight = trPr.find(f'{NS}trHeight')
            if trHeight is not None:
                val = int(trHeight.get(f'{NS}val', '500'))
                total += val / 20  # twips to pt
                continue
        total += 25  # default row height estimate
    return total


def _find_table_positions(doc):
    """
    找到表格在文档中的位置（在哪个段落之后）。
    返回 {paragraph_index: table_index} 的映射。
    """
    # 在 OOXML 中，表格 (w:tbl) 和段落 (w:p) 是 body 的直接子元素
    body = doc.element.body
    table_positions = {}
    para_count = -1
    table_idx = 0

    for child in body:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag == 'p':
            para_count += 1
        elif tag == 'tbl':
            # 表格紧跟在 para_count 段落之后
            table_positions[para_count] = table_idx
            table_idx += 1

    return table_positions


def estimate_page_numbers(doc):
    """估算文档中每个段落所在的页码（1-based）"""
    # 找 section break 位置
    section_breaks = []
    for i, p in enumerate(doc.paragraphs):
        pPr = p._element.find(f'{NS}pPr')
        if pPr is not None:
            sectPr = pPr.find(f'{NS}sectPr')
            if sectPr is not None:
                section_breaks.append(i)

    # 构建 section 范围
    sections = []
    start = 0
    for sb in section_breaks:
        sections.append((start, sb))
        start = sb + 1
    if start < len(doc.paragraphs):
        sections.append((start, len(doc.paragraphs) - 1))

    # 找表格位置
    table_positions = _find_table_positions(doc)

    # 估算每页
    page_map = {}
    current_page = 1
    content_start_page = None  # 正文 section 的起始绝对页码

    for sec_idx, (sec_start, sec_end) in enumerate(sections):
        if sec_idx >= len(doc.sections):
            break

        sec = doc.sections[sec_idx]
        usable_h_pt = (sec.page_height - sec.top_margin - sec.bottom_margin) / 12700
        usable_w_pt = (sec.page_width - sec.left_margin - sec.right_margin) / 12700

        # 分节符 = 新页
        if sec_idx > 0:
            current_page += 1
        current_y = 0

        # 最后一个 section 是正文，记录起始页
        if sec_idx == len(sections) - 1 and content_start_page is None:
            content_start_page = current_page

        for i in range(sec_start, sec_end + 1):
            # 段落高度
            h = _para_height_pt(doc.paragraphs[i], usable_w_pt)

            # 如果这个段落之后紧跟表格，加上表格高度
            if i in table_positions:
                t_idx = table_positions[i]
                if t_idx < len(doc.tables):
                    h += _table_height_pt(doc.tables[t_idx])

            # 检查翻页
            if current_y + h > usable_h_pt and current_y > 0:
                current_page += 1
                current_y = 0

            page_map[i] = current_page
            current_y += h

    # 转换为正文相对页码（从 1 开始）
    if content_start_page and content_start_page > 1:
        offset = content_start_page - 1
        page_map = {k: v - offset for k, v in page_map.items()}

    return page_map


def build_toc_mapping(doc):
    """构建目录项 → 正文标题的段落索引映射"""
    toc_entries = []
    for i in TOC_RANGE:
        runs_text = ''.join([r.text or '' for r in doc.paragraphs[i].runs])
        # 去掉 \t 以及后面的页码（可能是 ** 或数字）
        title = runs_text.split('\t')[0].strip()
        toc_entries.append((i, title))

    body_map = {}
    for toc_idx, toc_title in toc_entries:
        found = None

        for j in range(40, len(doc.paragraphs)):
            body_text = ''.join([r.text or '' for r in doc.paragraphs[j].runs]).strip()
            if body_text == toc_title:
                found = j
                break

        if found is None:
            clean = toc_title
            if clean.startswith('（') and '）' in clean:
                clean = clean[clean.index('）') + 1:]
            for j in range(40, len(doc.paragraphs)):
                body_text = ''.join([r.text or '' for r in doc.paragraphs[j].runs]).strip()
                if body_text == clean or body_text.endswith(clean):
                    found = j
                    break

        body_map[toc_idx] = {
            'title': toc_title,
            'body_idx': found,
        }

    return body_map


def update_toc_pages(filepath, body_map, page_map):
    """更新目录中的 ** 占位符"""
    doc = Document(filepath)

    updated = 0
    for toc_idx in TOC_RANGE:
        info = body_map[toc_idx]
        if info['body_idx'] is None:
            print(f"  跳过 toc[{toc_idx}] {info['title'][:30]}: 无映射")
            continue

        body_idx = info['body_idx']
        page = page_map.get(body_idx, '?')

        if page == '?':
            print(f"  跳过 toc[{toc_idx}] {info['title'][:30]}: 页码估算失败")
            continue

        para = doc.paragraphs[toc_idx]
        # 找到 tab 之后的 run（包含页码），替换为正确页码
        found_tab = False
        for run in para.runs:
            if found_tab:
                run.text = str(page)
                updated += 1
                break
            if '\t' in (run.text or ''):
                found_tab = True
                # 页码可能在同一个 run 中（如 "\t3"）或下一个 run
                if run.text.strip() and run.text.strip() != '\t':
                    # 页码和 tab 在同一个 run：保留 tab，替换后面的数字
                    run.text = '\t' + str(page)
                    updated += 1
                    break

    doc.save(filepath)
    print(f"  已更新 {updated} 个目录页码")
    return updated


def fix_file(filepath):
    """修复单个文件的目录页码"""
    print(f"\n处理: {os.path.basename(filepath)}")

    doc = Document(filepath)

    # 检测是否已有 TOC 域（真正的 Word 目录）
    if has_toc_field(doc):
        print("  文档已包含 TOC 域，请在 Word 中右键目录 → '更新域' 刷新页码")
        return True

    # 1. 估算页码
    print("  估算页码...")
    page_map = estimate_page_numbers(doc)

    # 2. 构建映射
    body_map = build_toc_mapping(doc)

    # 3. 打印结果
    print(f"  映射与页码:")
    for toc_idx in TOC_RANGE:
        info = body_map[toc_idx]
        body_idx = info['body_idx']
        page = page_map.get(body_idx, '?') if body_idx else '?'
        body_str = f'body[{body_idx}]' if body_idx else 'NOT FOUND'
        print(f'    toc[{toc_idx}] → {body_str} pg={page}  {info["title"][:25]}')

    # 4. 更新目录
    count = update_toc_pages(filepath, body_map, page_map)
    return count > 0


def main():
    files = [
        'AR眼镜AI应用开发课程标准(1).docx',
        '农业物联网技术课程标准.docx',
        '敏捷开发软件工程课程标准(1).docx',
    ]

    for fname in files:
        filepath = os.path.join(BASE_DIR, fname)
        if not os.path.exists(filepath):
            print(f"文件不存在: {filepath}")
            continue
        fix_file(filepath)


if __name__ == '__main__':
    main()
