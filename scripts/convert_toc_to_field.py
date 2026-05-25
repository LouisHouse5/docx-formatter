#!/usr/bin/env python3
"""
将静态文字目录转换为 Word 真正的 TOC 域目录。

静态目录：纯文字段落（标题 + \t + 页码），无法自动更新页码和跳转。
TOC 域目录：Word 原生目录功能，右键"更新域"即可自动刷新页码和超链接。

前提：正文标题需有 outlineLevel 属性（脚本自动添加）。

用法：
  python3 convert_toc_to_field.py 文档.docx
  python3 convert_toc_to_field.py --check 文档.docx        # 只检测不修改
  python3 convert_toc_to_field.py --levels "0-2" 文档.docx  # 指定大纲级别范围

注意：此脚本应在 copy_format_deep.py 之后运行（最后一步）。
"""
import re
import sys
import argparse
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from utils import has_toc_field, W_OUTLINELVL, W_FLDCHAR, W_INSTRTEXT

# 标题识别正则
H1_PATTERN = r'^[一二三四五六七八九十]+、'
H2_PATTERN = r'^（[一二三四五六七八九十]+）'

# TOC 域指令
TOC_INSTR = ' TOC \\o "1-2" \\h \\z \\u '


def _find_toc_title_idx(doc):
    """找到'目录'标题段落索引"""
    for i, p in enumerate(doc.paragraphs):
        text = ''.join(r.text or '' for r in p.runs).strip()
        if text == '目录':
            return i
    return None


def _detect_toc_range(doc, toc_title_idx):
    """
    自动检测目录条目范围（不依赖硬编码索引）。
    返回 (first_idx, last_idx) 或 None。
    """
    first = None
    last = None
    for i in range(toc_title_idx + 1, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        text = ''.join(r.text or '' for r in p.runs)

        # 检查是否包含制表符（目录条目的特征）
        if '\t' in text:
            if first is None:
                first = i
            last = i
        elif text.strip():
            # 非空且不含制表符 → 目录结束
            break
        else:
            # 空段落 → 目录结束
            break

    if first is not None and last is not None:
        return (first, last)
    return None


def _find_empty_para_after(doc, last_toc_idx):
    """找到目录条目之后的第一个空段落索引（用于放置 end 标记）"""
    for i in range(last_toc_idx + 1, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        text = ''.join(r.text or '' for r in p.runs).strip()
        if not text:
            return i
        # 如果遇到非空段落，也停止（可能在最后条目后没有空段落）
        break
    return None


def _get_heading_level(text):
    """根据文本模式判断标题级别"""
    if re.match(H1_PATTERN, text):
        return 0
    elif re.match(H2_PATTERN, text):
        return 1
    return None


def add_outline_levels(doc, toc_range=None):
    """
    为正文标题段落添加 outlineLevel 属性。
    跳过目录区域（由 toc_range 指定）。
    返回标记的段落数量。
    """
    skip_start = toc_range[0] if toc_range else -1
    skip_end = toc_range[1] if toc_range else -1
    count = 0

    for i, p in enumerate(doc.paragraphs):
        if skip_start <= i <= skip_end:
            continue

        text = ''.join(r.text or '' for r in p.runs).strip()
        level = _get_heading_level(text)
        if level is None:
            continue

        pPr = p._element.get_or_add_pPr()

        # 已有 outlineLevel 则跳过
        if pPr.find(qn('w:outlineLvl')) is not None:
            continue

        ol = OxmlElement('w:outlineLvl')
        ol.set(qn('w:val'), str(level))
        pPr.append(ol)
        count += 1

    return count


def _make_fldChar_run(char_type):
    """创建包含 fldChar 的 run 元素"""
    r = OxmlElement('w:r')
    fld = OxmlElement('w:fldChar')
    fld.set(qn('w:fldCharType'), char_type)
    r.append(fld)
    return r


def _make_instrText_run(instr):
    """创建包含 instrText 的 run 元素"""
    r = OxmlElement('w:r')
    instr_elem = OxmlElement('w:instrText')
    instr_elem.set(qn('xml:space'), 'preserve')
    instr_elem.text = instr
    r.append(instr_elem)
    return r


def insert_toc_field_markers(doc, first_toc_idx, last_toc_idx, toc_instr=None):
    """
    在目录区域嵌入 TOC 域标记（不增删段落）。

    - 在第一个目录条目段落前置插入 begin + instrText + separate runs
    - 在最后一个条目后的空段落中添加 end run
    """
    # 1. 在首个目录条目段落前置 field begin runs
    first_para = doc.paragraphs[first_toc_idx]._element
    pPr = first_para.find(qn('w:pPr'))

    begin_run = _make_fldChar_run('begin')
    instr_run = _make_instrText_run(toc_instr or TOC_INSTR)
    sep_run = _make_fldChar_run('separate')

    # 插入位置：pPr 之后（如果有的话），否则在最前面
    if pPr is not None:
        pPr.addnext(sep_run)
        pPr.addnext(instr_run)
        pPr.addnext(begin_run)
    else:
        first_para.insert(0, begin_run)
        first_para.insert(1, instr_run)
        first_para.insert(2, sep_run)

    # 2. 在空段落中添加 end 标记
    empty_idx = _find_empty_para_after(doc, last_toc_idx)
    if empty_idx is None:
        print("  警告：未找到目录条目后的空段落，无法放置 end 标记")
        return False

    end_para = doc.paragraphs[empty_idx]._element
    end_run = _make_fldChar_run('end')
    end_para.append(end_run)

    return True


def _build_toc_instr(levels='0-1'):
    """根据 levels 参数生成 TOC 域指令字符串"""
    # levels="0-1" 对应 \o "1-2"（Word 用 1-based 级别号）
    parts = levels.split('-')
    start = int(parts[0]) + 1
    end = int(parts[-1]) + 1
    return f' TOC \\o "{start}-{end}" \\h \\z \\u '


def convert_toc(filepath, levels='0-1', check_only=False):
    """
    主函数：将静态目录转换为 TOC 域。

    Args:
        filepath: 文档路径
        levels: 大纲级别范围（如 "0-1", "0-2"）
        check_only: 只检测不修改

    Returns:
        True 成功，False 失败
    """
    doc = Document(filepath)

    # 1. 检查是否已有 TOC 域
    if has_toc_field(doc):
        print("  文档已包含 TOC 域，跳过")
        return True

    # 2. 找到目录标题
    toc_title_idx = _find_toc_title_idx(doc)
    if toc_title_idx is None:
        print("  未找到'目录'标题段落")
        return False
    print(f"  目录标题位置：段落 [{toc_title_idx}]")

    # 3. 自动检测目录条目范围
    toc_range = _detect_toc_range(doc, toc_title_idx)
    if toc_range is None:
        print("  未检测到目录条目（含 \\t 的段落）")
        return False
    first_toc, last_toc = toc_range
    entry_count = last_toc - first_toc + 1
    print(f"  目录条目范围：段落 [{first_toc}] - [{last_toc}]（共 {entry_count} 条）")

    # 4. 检测模式：只报告
    if check_only:
        print(f"\n  === 检测报告 ===")
        print(f"  目录标题：段落 [{toc_title_idx}]")
        print(f"  条目数量：{entry_count}")
        print(f"  条目范围：[{first_toc}] - [{last_toc}]")

        # 统计正文标题
        h1_count = 0
        h2_count = 0
        for i, p in enumerate(doc.paragraphs):
            if first_toc <= i <= last_toc:
                continue
            text = ''.join(r.text or '' for r in p.runs).strip()
            level = _get_heading_level(text)
            if level == 0:
                h1_count += 1
            elif level == 1:
                h2_count += 1
        print(f"  正文一级标题：{h1_count} 个")
        print(f"  正文二级标题：{h2_count} 个")
        print(f"  TOC 指令：{_build_toc_instr(levels).strip()}")
        return True

    # 5. 为正文标题添加 outlineLevel
    outline_count = add_outline_levels(doc, toc_range=(first_toc, last_toc))
    print(f"  已为 {outline_count} 个正文标题添加 outlineLevel")

    # 6. 插入 TOC 域标记
    toc_instr = _build_toc_instr(levels)
    success = insert_toc_field_markers(doc, first_toc, last_toc, toc_instr)
    if not success:
        return False

    # 7. 保存
    doc.save(filepath)
    print(f"  ✓ 已保存：{filepath}")
    print(f"  ⚠️  请在 Word 中右键目录 → '更新域' → '更新整个目录'")
    return True


def main():
    parser = argparse.ArgumentParser(
        description='将静态文字目录转换为 Word 真正的 TOC 域目录'
    )
    parser.add_argument('file', help='目标 docx 文件路径')
    parser.add_argument('--check', action='store_true',
                        help='只检测不修改')
    parser.add_argument('--levels', default='0-1',
                        help='大纲级别范围（默认 "0-1"，对应 TOC \\o "1-2"）')
    args = parser.parse_args()

    print(f"文件: {args.file}")
    print("-" * 50)

    if args.check:
        convert_toc(args.file, levels=args.levels, check_only=True)
    else:
        success = convert_toc(args.file, levels=args.levels)
        if not success:
            sys.exit(1)


if __name__ == '__main__':
    main()
